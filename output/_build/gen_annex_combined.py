"""Combined Financial Annex: both project annexes (62.5 each) in one workbook, with formula links for the combined investor."""
import sys, re, json, copy, os; sys.path.insert(0, '.')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
import gen_annex as GA, packdata as PD
from engine2 import *
NAMES = ['Inputs', 'Source_Data', 'Project_CF', 'Sources_Uses', 'Option_A', 'Option_B', 'Scenarios', 'Sensitivity', 'Summary', 'Reconciliation']
pat = re.compile(r"(?<![A-Za-z0-9_'])('?)(" + '|'.join(NAMES) + r")('?)!")
AR = 'Arial'; F_N = Font(name=AR, size=10); F_B = Font(name=AR, size=10, bold=True); F_T = Font(name=AR, size=14, bold=True, color='1F3864'); F_IN = Font(name=AR, size=10, color='0000FF'); F_NOTE = Font(name=AR, size=9, italic=True, color='666666')
FILL_H = PatternFill('solid', fgColor='1F3864'); FILL_K = PatternFill('solid', fgColor='E8EEF7'); FILL_OK = PatternFill('solid', fgColor='E2EFDA'); F_H = Font(name=AR, size=10, bold=True, color='FFFFFF')
N1 = '#,##0.0;(#,##0.0);-'; PC = '0.0%'; X = '0.00"x"'; N2 = '#,##0.00;(#,##0.00);-'
def qc(i): return CL(3+i)
FIRST, LAST = qc(1), qc(60)
def make(Te, path, workdir):
    C = PD.build_combined(Te); wbn = openpyxl.Workbook(); wbn.remove(wbn.active); info = {}
    for k in ('GS', 'LA'):
        tmp = f'{workdir}/tmp_{k}.xlsx'; info[k] = GA.make_annex(k, Te, tmp); w = openpyxl.load_workbook(tmp)
        for n in NAMES:
            src = w[n]; dst = wbn.create_sheet(f'{k}_{n}'); dst.sheet_view.showGridLines = False
            for row in src.iter_rows():
                for c in row:
                    if c.value is None and not c.has_style: continue
                    v = c.value
                    if isinstance(v, str) and v.startswith('='): v = pat.sub(lambda m: f"{m.group(1)}{k}_{m.group(2)}{m.group(3)}!", v)
                    d = dst.cell(c.row, c.column, v)
                    if c.has_style: d.font = copy.copy(c.font); d.fill = copy.copy(c.fill); d.number_format = c.number_format; d.alignment = copy.copy(c.alignment); d.border = copy.copy(c.border)
            for col, dim in src.column_dimensions.items(): dst.column_dimensions[col].width = dim.width
            for r, dim in src.row_dimensions.items():
                if dim.height: dst.row_dimensions[r].height = dim.height
            dst.freeze_panes = src.freeze_panes
    G, L = info['GS'], info['LA']
    def trow(k):
        w = wbn[f'{k}_Inputs']
        return next(r for r in range(1, 40) if w.cell(r, 1).value == 'Investor ticket')
    TR = {k: trow(k) for k in ('GS', 'LA')}
    def sc(k, nm, key_row): r0 = info[k]['BLK'][('Scenarios', nm)]; return r0 + 7 + key_row   # raw0 coll1 land2 cons3 comm4 sga5 net6 cum7 a8 bpay9 binv10 cash11
    base_scn = 'Base case'; dn = 'Downside: sales +12 months, prices -10%'; up = 'Upside: prices +10%'; st = 'Delivery payments at handover'
    # ---- Combined_Summary
    ws = wbn.create_sheet('Combined_Summary'); ws.sheet_view.showGridLines = False; ws['A1'] = 'Combined investor: EGP 125 million, EGP 62.5 million in each project'; ws['A1'].font = F_T
    ws['A2'] = 'One investor, one option (A or B) applied to both projects. All cells are formulas linking to the two project annexes in this workbook.'; ws['A2'].font = F_NOTE
    ws.column_dimensions['A'].width = 62; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 14
    for i in range(1, 61): ws.column_dimensions[qc(i)].width = 9
    for j, h in enumerate(['Quarterly flows (EGP m)', 'Total', ''] + [''] * 60): c = ws.cell(4, 1+j, h); c.fill = FILL_H; c.font = F_H
    ws['A5'] = 'Month'; 
    for i in range(1, 61): ws[f'{qc(i)}5'] = f'=GS_Source_Data!{qc(i)}$6'; ws[f'{qc(i)}5'].font = F_N
    rows = [('a_base', 'Option A investor cash flow, base', lambda i: f"=GS_Option_A!{qc(i)}17+LA_Option_A!{qc(i)}17"),
            ('a_dn', 'Option A investor cash flow, downside', lambda i: f"=GS_Scenarios!{qc(i)}{sc('GS', dn, 8)}+LA_Scenarios!{qc(i)}{sc('LA', dn, 8)}"),
            ('a_up', 'Option A investor cash flow, upside', lambda i: f"=GS_Scenarios!{qc(i)}{sc('GS', up, 8)}+LA_Scenarios!{qc(i)}{sc('LA', up, 8)}"),
            ('j_a', 'Jalour net incremental cash flow, Option A', lambda i: f"=GS_Option_A!{qc(i)}21+LA_Option_A!{qc(i)}21"),
            ('b_inv', 'Option B investor cash flow (S1)', lambda i: f"=GS_Option_B!{qc(i)}14+LA_Option_B!{qc(i)}14"),
            ('j_b', 'Jalour cash flow, Option B (ticket less payouts)', lambda i: f"=GS_Option_B!{qc(i)}17+LA_Option_B!{qc(i)}17"),
            ('cash_b', 'Pooled project cash after tickets and payouts, base', lambda i: f"=GS_Option_B!{qc(i)}20+LA_Option_B!{qc(i)}20"),
            ('cash_dn', 'Pooled project cash after tickets and payouts, downside', lambda i: f"=GS_Scenarios!{qc(i)}{sc('GS', dn, 11)}+LA_Scenarios!{qc(i)}{sc('LA', dn, 11)}"),
            ('cash_st', 'Pooled project cash after tickets and payouts, delivery payments at handover', lambda i: f"=GS_Scenarios!{qc(i)}{sc('GS', st, 11)}+LA_Scenarios!{qc(i)}{sc('LA', st, 11)}")]
    R = {}
    for j, (key, lab, fn) in enumerate(rows):
        r = 6+j; R[key] = r; ws.cell(r, 1, lab).font = F_N
        for i in range(1, 61): c = ws[f'{qc(i)}{r}']; c.value = fn(i); c.number_format = N2; c.font = F_N
        ws.cell(r, 2, f'=SUM({FIRST}{r}:{LAST}{r})').number_format = N2
    rng = lambda key: f'{FIRST}{R[key]}:{LAST}{R[key]}'
    hdr_r = 17
    for j, h in enumerate(['Combined results', 'Value', 'Note']): c = ws.cell(hdr_r, 1+j, h); c.fill = FILL_H; c.font = F_H
    out = [('tick', 'Total investor ticket (EGP m)', f'=GS_Inputs!B{TR["GS"]}+LA_Inputs!B{TR["LA"]}', N1, 'EGP 62.5m in each project'), ('t_gs', 'Ticket in Green Square', f'=GS_Inputs!B{TR["GS"]}', N1, ''), ('t_la', "Ticket in L'avenir", f'=LA_Inputs!B{TR["LA"]}', N1, ''),
           ('face', 'Option A: units face value, both projects (EGP m)', '=GS_Option_A!B5+LA_Option_A!B5', N1, '2.0x the ticket at launch list'), ('val', 'Option A: units value at delivery list (EGP m)', '=GS_Option_A!B11+LA_Option_A!B11', N1, ''),
           ('a_irr', 'Option A investor IRR, base', f'=(1+IRR({rng("a_base")},0.05))^4-1', PC, ''), ('a_moic', 'Option A investor MOIC, base', f'=SUMIF({rng("a_base")},">0")/-SUMIF({rng("a_base")},"<0")', X, ''),
           ('a_dirr', 'Option A investor IRR, downside', f'=(1+IRR({rng("a_dn")},0.05))^4-1', PC, ''), ('a_dmoic', 'Option A investor MOIC, downside', f'=SUMIF({rng("a_dn")},">0")/-SUMIF({rng("a_dn")},"<0")', X, ''),
           ('a_uirr', 'Option A investor IRR, upside', f'=(1+IRR({rng("a_up")},0.05))^4-1', PC, ''), ('a_umoic', 'Option A investor MOIC, upside', f'=SUMIF({rng("a_up")},">0")/-SUMIF({rng("a_up")},"<0")', X, ''),
           ('j_irr', 'Jalour cost of capital, Option A', f'=(1+IRR({rng("j_a")},0.05))^4-1', PC, ''), ('j_npv', 'NPV cost to Jalour at 14%, Option A (EGP m)', f'=SUMPRODUCT({rng("j_a")},GS_Project_CF!{FIRST}{info["GS"]["P_"]["df"]}:{LAST}{info["GS"]["P_"]["df"]})', N1, ''),
           ('b_total', 'Option B: total return (EGP m)', '=GS_Option_B!B9+LA_Option_B!B9', N1, '2.3x the ticket'), ('b_irr', 'Option B investor IRR', f'=(1+IRR({rng("b_inv")},0.05))^4-1', PC, ''), ('b_moic', 'Option B investor MOIC', f'=SUMIF({rng("b_inv")},">0")/-SUMIF({rng("b_inv")},"<0")', X, ''),
           ('b_j', 'Jalour cost of capital, Option B', f'=(1+IRR({rng("j_b")},0.05))^4-1', PC, ''), ('b_cash', 'Lowest pooled cash after tickets and payouts, base (EGP m)', f'=MIN({rng("cash_b")})', N1, ''), ('b_cash_dn', 'Lowest pooled cash, downside (EGP m)', f'=MIN({rng("cash_dn")})', N1, ''), ('b_cash_st', 'Lowest pooled cash, delivery payments at handover (EGP m)', f'=MIN({rng("cash_st")})', N1, '')]
    M = {}
    for j, (key, lab, f_, fm, nt) in enumerate(out):
        r = hdr_r+1+j; ws.cell(r, 1, lab).font = F_N; c = ws.cell(r, 2, f_); c.number_format = fm; c.font = F_B; c.fill = FILL_K; ws.cell(r, 3, nt).font = F_NOTE; M[key] = f'Combined_Summary!$B${r}'
    # ---- Combined_Position
    wp = wbn.create_sheet('Combined_Position'); wp.sheet_view.showGridLines = False; wp['A1'] = "Green Square and L'avenir: combined cash position (sponsor level, before any investor instrument)"; wp['A1'].font = F_T
    wp.column_dimensions['A'].width = 58; wp.column_dimensions['B'].width = 16
    for i in range(1, 61): wp.column_dimensions[qc(i)].width = 9
    pcum = info['GS']['P_']['cum']
    wp['A5'] = 'Month'
    for i in range(1, 61): wp[f'{qc(i)}5'] = f'=GS_Source_Data!{qc(i)}$6'
    for r_, t_ in {6: 'Green Square cumulative net cash flow', 7: "L'avenir cumulative net cash flow", 8: 'Combined cumulative net cash flow', 9: 'Positive after the low (1/0)'}.items(): wp.cell(r_, 1, t_).font = F_B if r_ == 8 else F_N
    for i in range(1, 61):
        col = qc(i); wp[f'{col}6'] = f'=GS_Project_CF!{col}{pcum}'; wp[f'{col}7'] = f'=LA_Project_CF!{col}{pcum}'; wp[f'{col}8'] = f'={col}6+{col}7'; wp[f'{col}9'] = f'=IF(AND({col}8>0,{col}5>$B$12),1,0)'
        for r_ in (6, 7, 8): wp[f'{col}{r_}'].number_format = N2
    for r_, (t_, f_, fm) in {11: ('Lowest combined position', f'=MIN({FIRST}8:{LAST}8)', N1), 12: ('Month of the lowest combined position', f'=INDEX({FIRST}5:{LAST}5,MATCH(B11,{FIRST}8:{LAST}8,0))', '0'), 13: ('First month positive after the low', f'=INDEX({FIRST}5:{LAST}5,MATCH(1,{FIRST}9:{LAST}9,0))', '0')}.items():
        wp.cell(r_, 1, t_).font = F_B; c = wp.cell(r_, 2, f_); c.number_format = fm; c.font = F_B; c.fill = FILL_K
    M['cp_min'] = 'Combined_Position!$B$11'; M['cp_m'] = 'Combined_Position!$B$12'; M['cp_pos'] = 'Combined_Position!$B$13'
    row0 = 16
    for nm in PD.scen_defs():
        wp.cell(row0, 1, nm).font = F_B; wp.cell(row0, 1).fill = FILL_K; wp.cell(row0+1, 1, 'Green Square cumulative').font = F_N; wp.cell(row0+2, 1, "L'avenir cumulative").font = F_N; wp.cell(row0+3, 1, 'Combined cumulative').font = F_N
        for i in range(1, 61):
            col = qc(i); wp[f'{col}{row0+1}'] = f"=GS_Scenarios!{col}{sc('GS', nm, 7)}"; wp[f'{col}{row0+2}'] = f"=LA_Scenarios!{col}{sc('LA', nm, 7)}"; wp[f'{col}{row0+3}'] = f'={col}{row0+1}+{col}{row0+2}'
            for j_ in (1, 2, 3): wp[f'{col}{row0+j_}'].number_format = N2
        for j_, (t_, f_, fm) in enumerate([('Green Square peak shortage', f'=-MIN({FIRST}{row0+1}:{LAST}{row0+1})', N1), ("L'avenir peak shortage", f'=-MIN({FIRST}{row0+2}:{LAST}{row0+2})', N1), ('Combined peak shortage', f'=-MIN({FIRST}{row0+3}:{LAST}{row0+3})', N1), ('Month of combined peak', f'=INDEX({FIRST}$5:{LAST}$5,MATCH(MIN({FIRST}{row0+3}:{LAST}{row0+3}),{FIRST}{row0+3}:{LAST}{row0+3},0))', '0')]):
            wp.cell(row0+4+j_, 1, t_).font = F_B; c = wp.cell(row0+4+j_, 2, f_); c.number_format = fm; c.font = F_B; c.fill = FILL_K
        M['cp|' + nm] = f'Combined_Position!$B${row0+6}'; row0 += 10
    # ---- Combined_Reconciliation
    wr = wbn.create_sheet('Combined_Reconciliation'); wr.sheet_view.showGridLines = False; wr['A1'] = 'Combined reconciliation: all differences must read zero'; wr['A1'].font = F_T
    for j, h in enumerate(['Check', 'This workbook', 'Reference', 'Difference', 'Basis']): c = wr.cell(3, 1+j, h); c.fill = FILL_H; c.font = F_H
    A = C['A']; B1 = C['B']['S1']; sp = C['sponsor']
    checks = [('Project annex reconciliation, Green Square', '=IF(GS_Reconciliation!D%d=0,0,1)' % (info['GS']['last']+2), 0, 'Project annex'), ('Project annex reconciliation, L\'avenir', '=IF(LA_Reconciliation!D%d=0,0,1)' % (info['LA']['last']+2), 0, 'Project annex'),
              ('Combined tickets = 125', f'={M["tick"]}', 2*Te, 'Offer term'), ('Option A combined investor IRR', f'={M["a_irr"]}', A['inv']['Base']['irr'], 'Independent Python'), ('Option A combined investor MOIC', f'={M["a_moic"]}', A['inv']['Base']['moic'], 'Independent Python'),
              ('Option A combined downside IRR', f'={M["a_dirr"]}', A['inv']['Downside']['irr'], 'Independent Python'), ('Option A Jalour cost of capital', f'={M["j_irr"]}', A['jal']['irr'], 'Independent Python'), ('Option B combined investor IRR', f'={M["b_irr"]}', B1['Base']['irr'], 'Independent Python'),
              ('Option B lowest pooled cash, base', f'={M["b_cash"]}', B1['Base']['min_cash'], 'Independent Python'), ('Option B lowest pooled cash, downside', f'={M["b_cash_dn"]}', B1['Downside']['min_cash'], 'Independent Python'), ('Combined lowest cash position', f'={M["cp_min"]}', sp['comb_min'], 'Independent Python'),
              ('Combined downside peak shortage', f'={M["cp|" + dn]}', -sp['scen'][dn]['comb_peak'], 'Independent Python')]
    for j, (lab, a_, b_, bs) in enumerate(checks):
        r = 4+j; wr.cell(r, 1, lab).font = F_N; wr.cell(r, 2, a_).number_format = '#,##0.000000'; wr.cell(r, 3, b_).number_format = '#,##0.000000'; wr.cell(r, 3).font = F_IN
        c3 = wr.cell(r, 4, f'=ROUND(B{r}-C{r},5)'); c3.number_format = '0.00000'; c3.font = F_B; c3.fill = FILL_OK; wr.cell(r, 5, bs).font = F_NOTE
    last = 4+len(checks)-1
    wr.cell(last+2, 1, 'Sum of absolute differences').font = F_B; wr.cell(last+2, 4, f'=SUMPRODUCT(ABS(D4:D{last}))').number_format = '0.00000'; wr.cell(last+3, 1, 'Status').font = F_B; wr.cell(last+3, 4, f'=IF(D{last+2}=0,"ALL CHECKS ZERO","CHECK FAILED")').font = F_B
    for col, w in zip('ABCDE', (62, 18, 18, 14, 30)): wr.column_dimensions[col].width = w
    # ---- Cover
    cv = wbn.create_sheet('Cover', 0); cv.sheet_view.showGridLines = False; cv['A1'] = "Green Square and L'avenir: Financial Annex"; cv['A1'].font = Font(name=AR, size=18, bold=True, color='1F3864')
    cv['A2'] = 'Combined investor ticket EGP 125 million (EGP 62.5 million in each project)  |  Strictly private and confidential  |  Projections only, not guarantees'; cv['A2'].font = F_B
    notes = ['ALL INFORMATION IN THIS WORKBOOK IS CONFIDENTIAL AND SUBJECT TO THE NON-DISCLOSURE AGREEMENT. Do not copy, forward or disclose it.', 'Structure: sheets with the prefix GS_ are the Green Square annex; LA_ the L\'avenir annex (each with Inputs, Source_Data, Project_CF, Sources_Uses, Option_A, Option_B, Scenarios, Sensitivity, Summary, Reconciliation).', 'Combined_Summary, Combined_Position and Combined_Reconciliation add up the two projects for the combined investor.',
             'Colour code: blue = input or value transcribed from the source model; black = formula; green = link to another sheet. Units: EGP million unless stated.', 'The balance of project funding in each project is provided by Jalour sponsor equity, project collections and other capital sources.', 'Jalour may raise further capital at project or holding level, subject to the investor\'s stated rights.', 'Combined reconciliation status:']
    for j, t_ in enumerate(notes): cv.cell(4+j, 1, t_).font = F_B if j == 0 else F_N
    cv.cell(11, 2, f'=Combined_Reconciliation!D{last+3}').font = F_B; cv.cell(11, 2).fill = FILL_OK; cv.column_dimensions['A'].width = 150; cv.column_dimensions['B'].width = 22
    order = ['Cover', 'Combined_Summary', 'Combined_Position', 'Combined_Reconciliation'] + [f'{k}_{n}' for k in ('GS', 'LA') for n in NAMES]
    wbn._sheets = [wbn[n] for n in order]
    wbn.properties.title = "Financial Annex - Green Square and L'avenir - Combined ticket 125m"; wbn.properties.creator = 'Jalour Developments'; wbn.properties.subject = 'Confidential'
    wbn.save(path)
    return dict(M=M, info=info, C=C, last=last)
if __name__ == '__main__':
    r = make(62.5, sys.argv[1], os.path.dirname(sys.argv[1])); json.dump({'M': r['M']}, open(sys.argv[1] + '.map.json', 'w')); print('ok')
