import re, glob, os, json, openpyxl
from pptx import Presentation
OUT = '/home/user/Jalour_inv_pitch/output'
rows = []
for k in ('GS', 'LA'):
    for T in (50, 75, 100):
        rep = open(f'{OUT}/QA_Reports/QA_{k}_{T:03d}M.md').read(); f = f'{OUT}/{k}_{T:03d}M'
        tie = re.findall(r'\| (OK|CHECK) \|\n', rep + '\n'); n_ok = len(re.findall(r'\| OK \|$', rep, re.M)); n_chk = len(re.findall(r'\| CHECK \|$', rep, re.M))
        arith = re.findall(r'\| ([0-9.e+-]+) \| (OK|EXPLAIN) \|', rep); maxd = max(float(a) for a, _ in arith); nex = sum(1 for _, b in arith if b == 'EXPLAIN')
        hits = int(re.search(r'Prohibited-term hits: (\d+)', rep).group(1)); fails = int(re.search(r'Total failures requiring attention: (\d+)', rep).group(1))
        recon = 'ALL CHECKS ZERO' in rep
        slides = len(Presentation(f'{f}/Deck.pptx').slides)
        wb = openpyxl.load_workbook(f'{f}/Financial_Annex.xlsx'); nform = sum(1 for ws in wb for r in ws.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith('='))
        rows.append((f'{k}_{T:03d}M', slides, nform, n_ok, n_chk, maxd, nex, hits, 'yes' if recon else 'NO', fails))
md = ['# QA summary (INTERNAL)\n', 'Checks run on every pack: (1) number tie-out annex to deck, memo and DD memo; (2) independent recomputation of IRR, MOIC, NPV and coverage in pure Python versus the spreadsheet; (3) disclosure scan of every part of every file, including document properties, notes, hidden sheets and comments; (4) cross-document consistency; (5) visual render of decks and memos; (6) challenge pass with five sceptical-investor questions. Detailed reports: `QA_<pack>.md`.\n',
      '| Pack | Slides | Annex formulas | Tie-out OK | Tie-out CHECK | Max abs difference, independent vs spreadsheet | Differences to explain | Prohibited-term hits | Annex reconciliation all zero | Failures |', '|---|---|---|---|---|---|---|---|---|---|']
for r in rows: md.append(f'| {r[0]} | {r[1]} | {r[2]:,} | {r[3]} | {r[4]} | {r[5]:.1e} | {r[6]} | {r[7]} | {r[8]} | {r[9]} |')
md.append('\nNotes:\n- Six single-project packs (50, 75, 100M): each names only its own project; the disclosure scan forbids the other project name and every other prohibited term. The bare number "250" appears only as list prices (EGP thousand per m2) and in annex formulas, never as a programme total.\n- COMBINED_125M (EGP 62.5M in each project, the only 125M case) is checked separately in QA_COMBINED_125M.md: tie-out, independent recomputation, disclosure scan (both project names allowed there), confidentiality statements, NDA template content.\n- Every pack contains NDA_Template.docx; confidentiality statements appear on the cover and notice of every document and in the annex cover.')
open(f'{OUT}/QA_Reports/QA_Summary.md', 'w').write('\n'.join(md)); import re as _r; cf = open(f'{OUT}/QA_Reports/QA_COMBINED_125M.md').read(); md.append('\nCombined pack COMBINED_125M: failures ' + _r.search(r'Total failures requiring attention: (\d+)', cf).group(1) + '; annex reconciliation ' + ('ALL CHECKS ZERO' if 'ALL CHECKS ZERO' in cf else 'NOT ZERO'))
open(f'{OUT}/QA_Reports/QA_Summary.md', 'w').write('\n'.join(md)); print('\n'.join(md[:16]))
