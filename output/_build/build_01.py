import sys, json; sys.path.insert(0, '.')
import packdata as PD, numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
OUT = '/home/user/Jalour_inv_pitch/output'
HF = PatternFill('solid', fgColor='1F3864'); HFont = Font(bold=True, color='FFFFFF', name='Arial', size=10); BF = Font(name='Arial', size=10); TF = Font(name='Arial', size=13, bold=True); NOTE = Font(name='Arial', size=9, italic=True, color='555555'); FLAG = PatternFill('solid', fgColor='FCE4D6'); OKF = PatternFill('solid', fgColor='E2EFDA')
PCT, N1, X2 = '0.0%', '#,##0.0', '0.00"x"'
wb = openpyxl.Workbook(); wb.remove(wb.active)
def sheet(n, t, sub=None):
    ws = wb.create_sheet(n); ws['A1'] = t; ws['A1'].font = TF
    if sub: ws['A2'] = sub; ws['A2'].font = NOTE
    ws.sheet_view.showGridLines = False; return ws
def table(ws, r0, headers, rows, fmts=None, widths=None):
    for j, h in enumerate(headers):
        c = ws.cell(r0, 1+j, h); c.fill = HF; c.font = HFont; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(r0+1+i, 1+j, v); c.font = BF
            if fmts and fmts[j] and isinstance(v, (int, float)): c.number_format = fmts[j]
    if widths:
        for j, w in enumerate(widths): ws.column_dimensions[L(1+j)].width = w
    ws.row_dimensions[r0].height = 48; return r0+1+len(rows)
TK = (50, 75, 100, 125); DS = {(k, T): PD.build(k, T) for k in ('GS', 'LA') for T in TK}
ws = sheet('README', 'INTERNAL structuring analysis, revised after your decisions (Phase 2). Not for investors.', 'Source of truth: the two standalone project models (Green_Square.xlsx, Lavenir.xlsx; each project month 1 is its own sales launch). Results computed in Python (_build/engine2.py, packdata.py); the Financial Annexes recompute them with live formulas and tie to these values.')
for i, t in enumerate(['EGP million unless stated. NPV and IRR on quarter-end months, discount rate 14% (kept by Jalour). Investor funds at GS month -3 / LA month 0 (each project own timeline; combined view uses a common calendar with GS launch = month 1 and L\'avenir launch = month 7).',
        'Option A: units priced at the LAUNCH LIST (first sales tranche) with face value 2.0x the ticket, taken from the first two sales tranches. The 20% retained units are not used.',
        'Option B: 2.3x the ticket, 24-month grace, 12 equal quarterly instalments (recommended structure S1). S2 (share of collections) and S3 (hybrid) shown as alternatives.',
        'Nothing is ring-fenced: no reserve, escrow or collection account. The ticket is pooled project cash.',
        'Jalour bears the landlord 35% on investor units (no waiver). Jalour covers project funding gaps; no guarantee of return.',
        'Delivery payments before handover are contractual (kept as base); handover timing is shown as a stress.']): ws.cell(4+i, 1, t).font = BF
ws.column_dimensions['A'].width = 170
ws = sheet('Decisions', 'Your decisions and how they are applied')
table(ws, 3, ['Decision', 'Applied as'], [['Price appreciation 40-60%', 'Model list prices run launch to final: GS +33% retail / +48% offices; LA +44% / +50%. Investor resale at final list (also 40% and 60% flat cases).'], ['Jalour covers the funding gap', 'Stated as a sponsor funding undertaking (project gaps, not return). Gap sizes shown on Project_Scenarios.'], ['L\'avenir 595.3 basis', 'Used throughout; Summary range error recorded in the reconciliation.'],
    ['Ignore other projects in the source workbook', 'Not used anywhere. Excluded from all files.'], ['Do not touch retained units', 'Option A units come from the first two sales tranches; retained units are not used as units, collateral or payout source.'], ['Units: 2.0x; cash about 2.3x', 'Option A face 2.0x at launch list; Option B 2.3x.'], ['Nothing ring-fenced', 'Removed reserves, escrow, collection account and sweeps. Security = unit allocation / earmark, covenants, information rights, funding undertaking.'],
    ['Launch list basis', 'Units priced at 225/135 (GS) and 250/160 (LA) EGP thousand per m2.'], ['Delivery payment is contractual', 'Base case unchanged; stress retained.'], ['No landlord waiver', 'Jalour bears the 35% on investor units.'], ['Discount rate 14%', 'Kept.'], ['Avoid return guarantee', 'No guarantee; recommended substitutes in the DD memo.']], None, [44, 150])
ws = sheet('Project_Scenarios', 'Project scenarios (Jalour net cash flow before any investor instrument)')
rows = []
for k in ('GS', 'LA'):
    for nm, v in DS[(k, 100)]['proj'].items(): rows.append([PD.INFO[k]['name'], nm, v['net'], v['npv'], v['peak'], v['peak_m']])
table(ws, 3, ['Project', 'Scenario', 'Net cash flow', 'NPV @14%', 'Peak cumulative position', 'Peak month'], rows, [None, None, N1, N1, N1, '0'], [14, 44, 16, 14, 20, 10])
ws = sheet('Funding_Plan', 'Funding plan per ticket (nothing ring-fenced)', 'Pre-construction 10.5 is an assumption. The balance of project funding is provided by Jalour sponsor equity, project collections and other capital sources.')
rows = []
for (k, T), D in DS.items():
    f = D['fund']; u = f['uses']; rows.append([PD.INFO[k]['name'], T, u['dp'], u['guar_wc'], u['precon'], f['total'], f['applied'], f['balance'], f['excess_over_uses'], f['share_peak'], f['excess_over_peak'], 'EXCEEDS modelled peak' if f['excess_over_peak'] > 0 else '', f['stress_peak'], f['downside_peak']])
table(ws, 4, ['Project', 'Ticket', 'DP', 'Guarantee and working capital to peak', 'Pre-construction (assumption)', 'Total uses', 'Ticket applied', 'Balance', 'Ticket above uses', 'Ticket % of modelled peak', 'Ticket above modelled peak', 'Flag', 'Stress peak (delivery at handover)', 'Downside peak'], rows, [None, '0', N1, N1, N1, N1, N1, N1, N1, PCT, N1, None, N1, N1], [14, 8, 10, 18, 16, 12, 12, 12, 12, 14, 14, 20, 16, 14])
ws = sheet('OptA_Units', 'Option A: units at launch list, from the first two sales tranches', 'Jalour bears landlord 35%; landlord payments unchanged. Investor resale at final list less 3%, over four quarters from handover+3m.')
rows = []
for (k, T), D in DS.items():
    A = D['A']; ai = A['inv']; j = A['jal']
    rows.append([PD.INFO[k]['name'], T, A['face'], A['launch_value'], A['f'], A['retail_face'], A['office_face'], A['appr'], A['value_at_handover'], ai['Base']['irr'], ai['Base']['moic'], ai['Downside']['irr'], ai['Upside']['irr'], ai['Appreciation 40%']['irr'], ai['Appreciation 60%']['irr'], ai['Flat prices (0%)']['irr'], j['irr'], j['npv'], j['nominal'], -j['peak_ex_ticket'], j['prorata_irr'], j['prorata_npv']])
table(ws, 4, ['Project', 'Ticket', 'Face value', 'Launch tranche value', 'Share of tranches', 'Retail face', 'Office face', 'Appreciation x', 'Value at final list', 'Investor IRR base', 'MOIC base', 'IRR downside', 'IRR upside', 'IRR at +40%', 'IRR at +60%', 'IRR at 0%', 'Jalour cost of capital', 'NPV to Jalour @14%', 'Nominal cost', 'Peak shortage before ticket', 'Cost if pro rata all tranches', 'NPV if pro rata'], rows,
      [None, '0', N1, N1, PCT, N1, N1, '0.000', N1, PCT, X2, PCT, PCT, PCT, PCT, PCT, PCT, N1, N1, N1, PCT, N1], [14, 8, 10, 12, 10, 10, 10, 11, 12, 11, 9, 10, 10, 10, 10, 10, 12, 12, 10, 14, 14, 12])
ws = sheet('OptB_Cash', 'Option B: 2.3x, 24-month grace, 12 quarterly instalments', 'Strict coverage = quarterly net cash flow / instalment. Cash available = pooled cash (ticket plus cumulative net cash flow less payouts) plus the quarter.')
rows = []
for (k, T), D in DS.items():
    for S in ('S1', 'S2', 'S3'):
        r = D['B']['S'][S]['res']; b = r['Base']
        rows.append([PD.INFO[k]['name'], T, S + ' ' + D['B']['S'][S]['name'], b['irr'], b['moic'], b['jirr'], b['jnpv'], b['strict_min'], f"{b['strict_below']}/{b['n']}", b['liq_min'], b['min_cash'], r['Downside']['min_cash'], r['Stress']['min_cash'], r['Downside']['irr'], b['last']])
table(ws, 4, ['Project', 'Ticket', 'Structure', 'Investor IRR', 'MOIC', 'Jalour cost of capital', 'NPV to Jalour @14%', 'Lowest quarterly coverage (net CF)', 'Quarters below 1.5x', 'Lowest cash-available coverage', 'Lowest pooled cash, base', 'Lowest pooled cash, downside', 'Lowest pooled cash, delivery at handover', 'Investor IRR downside', 'Last payment month'], rows,
      [None, '0', None, PCT, X2, PCT, N1, '0.00', None, '0.00', N1, N1, N1, PCT, '0'], [14, 8, 40, 11, 8, 12, 12, 14, 10, 14, 12, 14, 16, 12, 10])
ws2 = sheet('Coverage_S1', 'Option B S1 payout coverage by quarter (ticket 100)')
r = 3
for k in ('GS', 'LA'):
    D = DS[(k, 100)]; S1 = D['B']['S']['S1']['res']['Base']; net = D['series']['net']; liq = dict(S1['liq'])
    ws2.cell(r, 1, PD.INFO[k]['name']).font = Font(name='Arial', bold=True)
    rows = [[m, net[(m+D['off'])//3-1], S1['pay'][(m+D['off'])//3-1], net[(m+D['off'])//3-1]/S1['pay'][(m+D['off'])//3-1], liq.get(m)] for m in D['B']['pay_months']]
    r = table(ws2, r+1, ['Month', 'Net cash flow', 'Instalment', 'Coverage (net CF)', 'Coverage (cash available)'], rows, ['0', N1, N1, '0.00', '0.00'], [10, 14, 12, 16, 20]) + 2
ws = sheet('Assumptions_Log', 'Assumptions (confirm)')
table(ws, 3, ['Ref', 'Assumption'], [['A1', 'Investor funds 100% at t0 (GS month -3, LA month 0); pooled with project cash, no restricted account.'], ['A2', 'Launch list = price list at the first sales tranche. Units come from the first two tranches (same price list).'], ['A3', 'Investor resale at final list price, 3% cost, four quarters from handover+3m.'], ['A4', 'Jalour cost of Option A = ticket received vs forgone launch-tranche collections plus commission saved; landlord payments unchanged.'],
    ['A5', 'Pre-construction cost 10.5 (30% of the 5% professional-fee line).'], ['A6', 'Downside: collections and commission +12 months, prices -10%; guarantee and construction timing unchanged.'], ['A7', 'Option B 2.3x per your instruction; no tuning to a target IRR.'], ['A8', 'Coverage "cash available" includes the ticket as pooled cash.']], None, [8, 150])

ws = sheet('Combined_Position', 'INTERNAL ONLY: combined Green Square + L\'avenir cash position (Jalour level). Never include in any investor pack.', 'Cumulative net cash flow before any investor instrument, EGP million, base case, from the edited source model.')
import engine as EN
gs, la = EN.series('GS')['cum'], EN.series('LA')['cum']
rows = [[int(m), float(gs[i]), float(la[i]), float(gs[i]+la[i])] for i, m in enumerate(EN.MONTHS[:43])]
r_end = table(ws, 4, ['Month', 'Green Square cumulative', "L'avenir cumulative", 'Combined cumulative'], rows, ['0', N1, N1, N1], [10, 22, 22, 22])
mn = min(rows, key=lambda x: x[3]); ws.cell(r_end+1, 1, f'Lowest combined position: {mn[3]:.1f} at month {mn[0]} (Green Square alone {min(r[1] for r in rows):.1f}; L\'avenir alone {min(r[2] for r in rows):.1f}).').font = BF
wb.save(f'{OUT}/01_Structuring_Analysis.xlsx')
# summary md
g = DS[('GS', 100)]; l = DS[('LA', 100)]
def P1(x): return f'{x*100:.1f}%'
md = f"""# 01 Structuring Summary (INTERNAL, revised for your decisions)

Projection only. EGP million. Detail: `01_Structuring_Analysis.xlsx`; every pack Financial Annex recomputes these numbers with live formulas and reconciles to the source model (GS net 375.4 / NPV 229.9; LA net 595.3 / NPV 350.6, each at its own month 0).

## Applied from your decisions
Retained units untouched; units at **launch list** (GS retail 225 / offices 135; LA 250 / 160 EGP thousand per m2) with face value 2.0x; cash 2.3x; nothing ring-fenced; Jalour bears the landlord 35%; Jalour covers project funding gaps; no return guarantee; discount rate 14%; other projects in the source workbook ignored.

## Option A: units (2.0x at launch list, from the first two sales tranches)
| Per ticket 100 | Green Square | L'avenir |
|---|---|---|
| Units face value / value at final list | 200 / {g['A']['value_at_handover']:.0f} | 200 / {l['A']['value_at_handover']:.0f} |
| Share of launch tranches (tranche value {g['A']['launch_value']:.0f} / {l['A']['launch_value']:.0f}) | {P1(g['A']['f'])} | {P1(l['A']['f'])} |
| Investor IRR / MOIC: base | {P1(g['A']['inv']['Base']['irr'])} / {g['A']['inv']['Base']['moic']:.2f}x | {P1(l['A']['inv']['Base']['irr'])} / {l['A']['inv']['Base']['moic']:.2f}x |
| Downside (resale +12m, price -10%) / upside (+10%) | {P1(g['A']['inv']['Downside']['irr'])} / {P1(g['A']['inv']['Upside']['irr'])} | {P1(l['A']['inv']['Downside']['irr'])} / {P1(l['A']['inv']['Upside']['irr'])} |
| Appreciation +40% / +60% / none | {P1(g['A']['inv']['Appreciation 40%']['irr'])} / {P1(g['A']['inv']['Appreciation 60%']['irr'])} / {P1(g['A']['inv']['Flat prices (0%)']['irr'])} | {P1(l['A']['inv']['Appreciation 40%']['irr'])} / {P1(l['A']['inv']['Appreciation 60%']['irr'])} / {P1(l['A']['inv']['Flat prices (0%)']['irr'])} |
| Jalour cost of capital / NPV cost at 14% | {P1(g['A']['jal']['irr'])} / {g['A']['jal']['npv']:.1f} | {P1(l['A']['jal']['irr'])} / {l['A']['jal']['npv']:.1f} |
| Peak shortage before ticket (base 102.5 / 102.0) | {-g['A']['jal']['peak_ex_ticket']:.1f} | {-l['A']['jal']['peak_ex_ticket']:.1f} |

Costs scale linearly with the ticket (NPV cost at 50 / 75 / 125: GS {DS[('GS',50)]['A']['jal']['npv']:.1f} / {DS[('GS',75)]['A']['jal']['npv']:.1f} / {DS[('GS',125)]['A']['jal']['npv']:.1f}; LA {DS[('LA',50)]['A']['jal']['npv']:.1f} / {DS[('LA',75)]['A']['jal']['npv']:.1f} / {DS[('LA',125)]['A']['jal']['npv']:.1f}). The units need 17% to 41% of the first two tranches; retail-only breaks above 116M (GS) so a pro-rata mix (about 32% retail / 68% offices) is used. Taking the same area pro rata from all tranches costs more (GS {P1(g['A']['jal']['prorata_irr'])}, LA {P1(l['A']['jal']['prorata_irr'])}). **Cost to Jalour is a real transfer:** the investor receives launch-priced area that appreciates 43% to 49% while Jalour forgoes the launch-tranche collections; the investor's units also compete with Jalour's later sales.

## Option B: cash (2.3x, 24m grace, 12 quarterly instalments)
Investor IRR **{P1(g['B']['S']['S1']['res']['Base']['irr'])}** (both projects; same timing from funding), MOIC 2.30x; Jalour cost equals the investor IRR. Option A and B are close for Green Square ({P1(g['A']['inv']['Base']['irr'])} vs {P1(g['B']['S']['S1']['res']['Base']['irr'])}); Option A is richer for L'avenir ({P1(l['A']['inv']['Base']['irr'])}).
| Ticket 100 | GS | LA |
|---|---|---|
| S1 fixed: quarters with net-cash-flow coverage below 1.5x | {g['B']['S']['S1']['res']['Base']['strict_below']} of 12 | {l['B']['S']['S1']['res']['Base']['strict_below']} of 12 |
| S1 lowest cash-available coverage | {g['B']['S']['S1']['res']['Base']['liq_min']:.1f}x | {l['B']['S']['S1']['res']['Base']['liq_min']:.1f}x |
| S2 % of collections (floor 0.4x, cap 1.5x): share / IRR | {P1(g['B']['params']['S2_pct'])} / {P1(g['B']['S']['S2']['res']['Base']['irr'])} | {P1(l['B']['params']['S2_pct'])} / {P1(l['B']['S']['S2']['res']['Base']['irr'])} |
| S3 hybrid (1.8x fixed + share, cap 2.8x): share / IRR | {P1(g['B']['params']['S3_pct'])} / {P1(g['B']['S']['S3']['res']['Base']['irr'])} | {P1(l['B']['params']['S3_pct'])} / {P1(l['B']['S']['S3']['res']['Base']['irr'])} |
| Lowest pooled cash after payouts: base / downside / stress | {g['B']['S']['S1']['res']['Base']['min_cash']:.0f} / {g['B']['S']['S1']['res']['Downside']['min_cash']:.0f} / {g['B']['S']['S1']['res']['Stress']['min_cash']:.0f} | {l['B']['S']['S1']['res']['Base']['min_cash']:.0f} / {l['B']['S']['S1']['res']['Downside']['min_cash']:.0f} / {l['B']['S']['S1']['res']['Stress']['min_cash']:.0f} |

**Recommend S1 (fixed 2.3x, equal instalments).** It matches your structure and is the simplest to close. **No structure passes a 1.5x test on quarterly net cash flow in every quarter** (the construction peak is net negative: GS months 33 to 42, LA 36 to 42), and with no ring-fence nothing can cure that. I disclose it, show pooled-cash coverage (at least {g['B']['S']['S1']['res']['Base']['liq_min']:.1f}x GS, {l['B']['S']['S1']['res']['Base']['liq_min']:.1f}x LA in base) and rely on Jalour's funding undertaking. At 125M GS the pooled-cash coverage is only {DS[('GS',125)]['B']['S']['S1']['res']['Base']['liq_min']:.2f}x in one quarter.
Percent-of-collections share per ticket (S2): GS {', '.join(P1(DS[('GS',T)]['B']['params']['S2_pct']) for T in TK)}; LA {', '.join(P1(DS[('LA',T)]['B']['params']['S2_pct']) for T in TK)} at 50 / 75 / 100 / 125.

## What the funding gap looks like (you will fund this)
Delivery payments at handover: GS {g['fund']['stress_peak']:.0f}, LA {l['fund']['stress_peak']:.0f}. Sales +12 months and prices -10%: GS {g['fund']['downside_peak']:.0f}, LA {l['fund']['downside_peak']:.0f}. Both exceed any ticket several times over. Ticket 125 exceeds the modelled peak by 22.5 / 23.0 (general project liquidity, not restricted).

## No guarantee: recommendation (decision 5)
Do not guarantee the return. Offer instead: payment priority covenant (no shareholder distributions while an instalment is overdue), earmark of unsold launch-tranche or later units worth at least the ticket with substitution, negative pledge, monthly reporting and audit rights, and a sponsor funding undertaking limited to project funding gaps and completion. Optional: promissory notes, independent engineer, step-in right.

## Items resolved or changed after your answers
1. Funding gaps: to be covered from Jalour's cash reserves and cash from other projects it is constructing. The packs say so, and say plainly that this capacity is not ring-fenced and depends on those projects (GS downside gap {g['fund']['downside_peak']:.0f}, LA {l['fund']['downside_peak']:.0f}). Form and any cap still to be documented.
2. Separate project account: treated as not applicable to these plots (Jalour's position); basis goes in the data room.
3. Prices compared with delivery prices: Option A and the price pages now use the delivery list (last list price in the plan, unchanged to handover). GS offices +{g['pricecmp']['office_g']*100:.0f}% and retail +{g['pricecmp']['retail_g']*100:.0f}% from launch; LA +{l['pricecmp']['office_g']*100:.0f}% and +{l['pricecmp']['retail_g']*100:.0f}%. Against today's office asking prices the delivery list is +{g['pricecmp']['off_vs_ask']*100:.0f}% (GS, about {g['pricecmp']['off_cagr']*100:.0f}% a year) and +{l['pricecmp']['off_vs_ask']*100:.0f}% (LA, about {l['pricecmp']['off_cagr']*100:.0f}% a year).
4. Market fact base: sent to you for review before verification.
5. Sponsor, permit, title and landlord-consent documents: marked as provided after the investor signs an NDA.
"""
md += "\n\n**Note (125M).** A 125M ticket means one investor placing 62.5M in each project (pack COMBINED_125M). The single-project 125 columns above are analysis only; no single-project 125M pack exists.\n"
open(f'{OUT}/01_Structuring_Summary.md', 'w').write(md); print('ok')
