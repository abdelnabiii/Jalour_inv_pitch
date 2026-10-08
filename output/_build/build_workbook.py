import sys, json; sys.path.insert(0, '.')
from engine import *
from run_analysis import R, TICKETS, M_B, FACE, PRECON
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
wb = openpyxl.Workbook()
HF = PatternFill('solid', fgColor='1F3864'); HFont = Font(bold=True, color='FFFFFF', name='Arial', size=10); BF = Font(name='Arial', size=10)
INP = Font(name='Arial', size=10, color='0000FF'); TF = Font(name='Arial', size=13, bold=True); NOTE = Font(name='Arial', size=9, italic=True, color='555555')
FLAG = PatternFill('solid', fgColor='FCE4D6'); OK = PatternFill('solid', fgColor='E2EFDA')
def sheet(name, title, sub=None):
    ws = wb.create_sheet(name); ws['A1'] = title; ws['A1'].font = TF
    if sub: ws['A2'] = sub; ws['A2'].font = NOTE
    ws.sheet_view.showGridLines = False; return ws
def table(ws, r0, headers, rows, fmts=None, widths=None, flagcol=None):
    for j, h in enumerate(headers):
        c = ws.cell(r0, 1+j, h); c.fill = HF; c.font = HFont; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(r0+1+i, 1+j, v); c.font = BF
            if fmts and fmts[j] and isinstance(v, (int, float)): c.number_format = fmts[j]
    if widths:
        for j, w in enumerate(widths): ws.column_dimensions[L(1+j)].width = w
    return r0+1+len(rows)
PCT, N1, N2, X2 = '0.0%', '#,##0.0', '#,##0.00', '0.00"x"'
wb.remove(wb.active)
# ---- README
ws = sheet('README', 'INTERNAL - Structuring Analysis (Phase 2). Not for distribution to investors.',
           'Contains both projects and combined figures. Source of truth: ALAHLY-SABBOUR MOSTAKBAL CITY -V2.xlsx (cached values, extracted by extract_model.py).')
lines = ['Units: EGP million unless stated. Time axis: quarter-end months (M3, M6, ...); NPV/IRR use t = month/12, annual compounding. Model discount rate 14% p.a. (Green Square sheet C72, labelled "Cost Of Capital").',
 'Base_CF reproduces the model: GS net 375.4 / NPV 215.3; LA net 595.3 / NPV 307.5 (check cells on that sheet).',
 'Scenario, Option A and Option B results were computed in Python (engine.py, run_analysis.py) and written here as values. The Phase 3 Financial Annex will be formula-driven.',
 'Tags: MODEL = taken from the source model. ASSUMPTION = mine, needs your confirmation (see Assumptions_Log). Blue font = input.',
 'Investor funding date t0: GS M3 (down payment + first instalment due), LA M12 (down payment due). Full ticket assumed funded at t0.',
 'Option A face value is measured at the MODEL AVERAGE REALISED PRICE (not launch list price). See OptA_Sources for the launch-price alternative.',
 'Option B: 24-month grace from t0, then 12 quarterly payout slots; total return multiple 1.75x of ticket (chosen for IRR parity with Option A, see OptB_Summary).']
for i, t in enumerate(lines): ws.cell(4+i, 1, t).font = BF
ws.column_dimensions['A'].width = 170
# ---- Base_CF with formulas
ws = sheet('Base_CF', 'Base case quarterly cash flows (MODEL) with live net, cumulative and PV formulas', 'Collections, landlord payments and costs extracted from Net Cash Flow -With DP rows 4/5, 11/12, 32/33.')
ws['A3'] = 'Discount rate'; ws['B3'] = DISC; ws['B3'].font = INP; ws['B3'].number_format = PCT
r = 5
for k in P:
    s = series(k); ws.cell(r, 1, P[k]['name']).font = Font(name='Arial', bold=True, size=11)
    hdr = ['Month', 't (yrs)', 'Collections (sold 80%)', 'Landlord payments (incl. DP)', 'Construction', 'Sales commission', 'SG&A & others', 'Net cash flow', 'Cumulative', 'Discount factor', 'PV of net']
    for j, h in enumerate(hdr):
        c = ws.cell(r+1, 1+j, h); c.fill = HF; c.font = HFont; c.alignment = Alignment(wrap_text=True, horizontal='center')
    first = r+2
    for i in range(43):
        rr = first+i; q = i
        vals = [int(MONTHS[i]), f'=A{rr}/12', float(s['coll'][i]), float(s['land'][i]), float(s['cons'][i]), float(s['comm'][i]), float(s['sga'][i])]
        for j, v in enumerate(vals):
            c = ws.cell(rr, 1+j, v); c.font = INP if j >= 2 else BF; c.number_format = N2 if j >= 2 else '0.00'
        ws.cell(rr, 8, f'=C{rr}-D{rr}-E{rr}-F{rr}-G{rr}').number_format = N2
        ws.cell(rr, 9, f'=H{rr}' if i == 0 else f'=I{rr-1}+H{rr}').number_format = N2
        ws.cell(rr, 10, f'=1/(1+$B$3)^B{rr}').number_format = '0.0000'
        ws.cell(rr, 11, f'=H{rr}*J{rr}').number_format = N2
        for j in (8, 9, 10, 11): ws.cell(rr, j).font = BF
    last = first+42; tr = last+1
    ws.cell(tr, 1, 'Total / check').font = Font(name='Arial', bold=True)
    for j in (3, 4, 5, 6, 7, 8, 11): ws.cell(tr, j, f'=SUM({L(j)}{first}:{L(j)}{last})').number_format = N2
    ws.cell(tr+1, 1, 'Peak cumulative shortage'); ws.cell(tr+1, 9, f'=MIN(I{first}:I{last})').number_format = N2
    exp_net = sum(D[k]['net']); exp_npv = {'GS': 215.31139254892392, 'LA': 307.5323710915294}[k]
    ws.cell(tr+2, 1, 'Model net / NPV (hardcoded from model)'); ws.cell(tr+2, 8, exp_net).number_format = N2; ws.cell(tr+2, 11, exp_npv).number_format = N2
    ws.cell(tr+3, 1, 'Check (must be 0.00)'); ws.cell(tr+3, 8, f'=ROUND(H{tr}-H{tr+2},2)'); ws.cell(tr+3, 11, f'=ROUND(K{tr}-K{tr+2},2)')
    for j in (8, 11): ws.cell(tr+3, j).fill = OK
    r = tr+6
for j, w in enumerate([10, 8, 16, 16, 14, 14, 14, 14, 14, 12, 14]): ws.column_dimensions[L(1+j)].width = w
# ---- Model_Tieout
ws = sheet('Model_Tieout', 'Headline figures: your prompt vs source model', 'Highlighted rows = conflict with the prompt. Stopped and listed per your instruction.')
rows = [
 ['Land area (m2)', 14470, 14470, 14910, 14910, 'OK'], ['Total BUA (m2)', 13023, 13023, 13419, 13419, 'OK'],
 ['Sales start / construction start / delivery (month)', '7/13/48', '7/13/48', '13/19/54', '13/19/54', 'OK'],
 ['Total sales 100% (EGP m)', 2816, 2816.2, 3342, 3342.2, 'OK'], ['Total sales 80% (EGP m)', 2253, 2253.0, 2674, 2673.7, 'OK'],
 ['Landlord min guarantee / DP (EGP m)', '800 / 80', '800 / 80 (Offers DP)', '850 / 85', '850 / 85 (Offers DP)', 'OK - Summary!B3:B4 show 200 but are labels only, not used in any calculation'],
 ['Jalour cash inflow (EGP m)', 1267, 1267.3, 1504, 1504.0, 'OK'],
 ['Jalour cash outflow (EGP m)', 892, 891.9, 862, 908.7, 'CONFLICT (LA): prompt/Summary 862; cash-flow engine 908.7'],
 ['Jalour net cash flow (EGP m)', 375, 375.4, 642, 595.3, 'CONFLICT (LA): prompt 642 (L\'avenir Summary C23); engine and Summary!D4 595.3'],
 ['NPV with DP @14% (EGP m)', 215.3, 215.3, 307.5, 307.5, 'OK (NPV uses the 908.7 cost, i.e. 595.3 basis)'],
 ['NPV without DP (EGP m)', 241.2, 241.2, 328.6, 328.6, 'OK'], ['Peak cumulative shortage (EGP m)', 102.5, 102.5, 102.0, 102.0, 'OK (GS at M6, LA at M30)'],
 ['Construction cost (EGP m)', 697, 697.2, 651, 697.2, 'CONFLICT (LA): 651 is a truncated SUM (Cash Out Detail!C29 =SUM(AB29:CC29) omits V:AA = 46.4)'],
 ['Sales commission (EGP m)', 90, 90.1, 107, 106.9, 'OK'], ['SG&A (EGP m)', 105, 104.6, 105, 104.6, 'OK - identical because both = 15% x 697.2 (see reconciliation)']]
table(ws, 4, ['Item', 'GS prompt', 'GS model', 'LA prompt', 'LA model', 'Status'], rows, widths=[48, 14, 18, 14, 18, 100])
for i, row in enumerate(rows):
    if row[5].startswith('CONFLICT'):
        for j in range(6): ws.cell(5+i, 1+j).fill = FLAG
# ---- Project_Scenarios
ws = sheet('Project_Scenarios', 'Project-level scenarios (Jalour net cash flow, before any investor instrument)', 'ASSUMPTION definitions: sales delay shifts all collections and commission; guarantee schedule and construction timing unchanged; price % scales collections and landlord excess.')
rows = []
for k in P:
    for nm, v in R['proj'][k].items(): rows.append([P[k]['name'], nm, v['net'], v['npv'], v['peak'], v['peak_m']])
table(ws, 4, ['Project', 'Scenario', 'Net cash flow (EGP m)', 'NPV @14% (EGP m)', 'Peak cumulative cash position (EGP m)', 'Peak month'], rows, [None, None, N1, N1, N1, '0'], [14, 44, 18, 16, 22, 10])
ws.cell(6+len(rows), 1, 'Reading: the model collects delivery payments (20-25% of price) at M30-M36 (GS) and M39-M45 (LA), i.e. 12-18 and 9-15 months BEFORE the model\'s own handover month (48 / 54). The "Delivery payments at handover" row moves them to handover.').font = NOTE
# ---- Funding_Plan
ws = sheet('Funding_Plan', 'Funding plan per ticket (per project)', 'Uses: DP and guarantee/working capital to modelled peak are MODEL; pre-construction is an ASSUMPTION (30% of the 5% professional-fee line, 10.5m, pulled ahead of construction start).')
rows = []
for k in P:
    f = R['fund'][k]; u = f['uses']
    for T in TICKETS:
        x = f['rows'][T]
        rows.append([P[k]['name'], T, u['down_payment'], u['guarantee_wc_to_peak'], u['precon_assumption'], f['total_uses'], min(T, f['total_uses']), x['balance'], x['buffer'], x['share_model_peak'], x['excess_over_model_peak'], 'EXCEEDS modelled peak' if T > f['model_peak'] else ''])
table(ws, 4, ['Project', 'Ticket', 'Use: down payment', 'Use: guarantee instalments & working capital to peak', 'Use: design/permits/pre-construction (assumption)', 'Total uses at peak', 'Ticket applied to uses', 'Balance: sponsor equity, collections and other sources', 'Ticket above uses = buffer / reserve', 'Ticket as % of modelled peak', 'Ticket above modelled peak', 'Flag'],
      rows, [None, '0', N1, N1, N1, N1, N1, N1, N1, PCT, N1, None], [14, 8, 14, 22, 22, 14, 14, 24, 18, 14, 14, 22])
ws.cell(6+len(rows), 1, 'Stress: if delivery payments arrive at handover, peak shortage is GS 269 / LA 244, above every ticket. Downside (sales +12m, price -10%): GS 441 / LA 849. No ticket alone covers these; see summary.').font = NOTE
# ---- OptA_Sources
ws = sheet('OptA_Sources', 'Option A: which inventory funds the units?', 'Pools at the model average realised price. Retained pool = 20% of units not in the model\'s sales plan and not in its cash flows.')
rows = []
for k in P:
    p = P[k]; f = R['facts'][k]
    ratio = (p['ret_comm']+p['ret_adm'])/(p['ret_comm']/(f['avg_comm']/f['launch_comm'])+p['ret_adm']/(f['avg_adm']/f['launch_adm']))
    for T in TICKETS:
        x = R['A']['pool'][f'{k}|{T}']
        rows.append([p['name'], T, x['face'], p['retained_val'], x['pct_retained'], p['ret_comm'], p['ret_adm'], x['pct_sold'], x['pct_list100'], ratio, x['face']*ratio])
table(ws, 4, ['Project', 'Ticket', 'Face value of units (2.0x)', 'Retained pool value', 'Face as % of retained pool', 'Retained pool: retail', 'Retained pool: offices', 'Face as % of sold programme (80%)', 'Face as % of total sales value (100%)', 'Avg price / launch price (blended, pool-weighted)', 'Face value if measured at launch list price'],
      rows, [None, '0', N1, N1, PCT, N1, N1, PCT, PCT, '0.00', N1], [14, 8, 16, 16, 16, 16, 16, 18, 18, 22, 22])
ws.cell(6+len(rows), 1, 'Retail-only allocation exceeds the retail retained pool above a ticket of about 116m (GS) / 135m (LA): a retail/office mix is needed at 125m. Offices-only fits all tickets.').font = NOTE
# ---- OptA_Investor
ws = sheet('OptA_Investor', 'Option A: investor returns (projection, 2.0x face at average modelled price, resold over 4 quarters from handover+3m, 3% resale friction)', 'Downside: resale 12 months later and 10% lower price. Upside: 10% higher price. Returns are the same per EGP at every ticket because the instrument is linear.')
rows = []
for k in P:
    for T in TICKETS:
        row = [P[k]['name'], T, FACE*T]
        for sc in ('Base', 'Downside', 'Upside'): x = R['A']['inv'][f'{k}|{T}|{sc}']; row += [x['irr'], x['moic']]
        rows.append(row)
table(ws, 4, ['Project', 'Ticket', 'Units face value', 'Base IRR', 'Base MOIC', 'Downside IRR', 'Downside MOIC', 'Upside IRR', 'Upside MOIC'], rows, [None, '0', N1, PCT, X2, PCT, X2, PCT, X2], [14, 8, 16, 12, 12, 14, 14, 12, 12])
r0 = 6+len(rows); ws.cell(r0, 1, 'Face multiple sensitivity (ticket 100m): investor IRR / MOIC').font = Font(name='Arial', bold=True)
rows2 = [[P[k]['name'], float(fx), R['A']['face_sens'][f'{k}|{fx}']['irr'], R['A']['face_sens'][f'{k}|{fx}']['moic']] for k in P for fx in (1.5, 1.75, 2.0, 2.25, 2.5, 2.67, 3.0)]
r1 = table(ws, r0+1, ['Project', 'Face multiple', 'IRR', 'MOIC'], rows2, [None, X2, PCT, X2])
ws.cell(r1+1, 1, 'Early assignment 24 months after funding at a discount to face (ticket 100m, same for both projects): investor IRR').font = Font(name='Arial', bold=True)
rows3 = [[float(d), R['A']['early_exit'][f'GS|{d}']['irr'], R['A']['early_exit'][f'GS|{d}']['moic']] for d in (0.0, 0.1, 0.2, 0.3)]
table(ws, r1+2, ['Discount to face', 'IRR', 'MOIC'], rows3, [PCT, PCT, X2])
# ---- OptA_Jalour_Cost
ws = sheet('OptA_Jalour_Cost', 'Option A: what the 2.0x units cost Jalour', 'Cost rate = discount rate at which +ticket now equals -value of units given up later (retained: forgone proceeds net of 4% commission; sold: lost collections plus saved commission). "Bears" = landlord 35% still paid in full (as in the model). "Waives" = landlord forgoes its 35% on investor units (needs Al Ahly Sabbour consent).')
rows = []
for k in P:
    for T in TICKETS:
        for src in ('retained', 'sold'):
            for ll in ('bears', 'waives'):
                x = R['A']['jal'][f'{k}|{T}|{src}|{ll}']
                rows.append([P[k]['name'], T, 'Retained 20% pool' if src == 'retained' else 'Sold programme (pre-sale discount)', 'Jalour bears landlord share' if ll == 'bears' else 'Landlord waives share', x['nominal'], x['npv14'], x['cost_rate'], x['peak_ex_ticket'], x['npv_project_after']])
table(ws, 4, ['Project', 'Ticket', 'Unit source', 'Landlord 35% on investor units', 'Nominal net cost (EGP m)', 'NPV @14% of the deal to Jalour (EGP m)', 'Jalour cost of capital (IRR)', 'Peak cash position before ticket (EGP m)', 'Project NPV after deal (EGP m)'],
      rows, [None, '0', None, None, N1, N1, PCT, N1, N1], [14, 8, 34, 28, 16, 18, 16, 22, 18])
r1 = 6+len(rows); ws.cell(r1, 1, 'If Jalour would otherwise have held the retained units longer (ticket 100m, retained, Jalour bears): cost of capital').font = Font(name='Arial', bold=True)
rows2 = [[P[k]['name'], f'sold at handover+3m +{d}m', R['A']['jal'][f'{k}|100|retained|bears|hold+{d}']['cost_rate']] for k in P for d in (0, 12, 24)]
table(ws, r1+1, ['Project', 'Retained-unit sale timing forgone', 'Cost of capital'], rows2, [None, None, PCT])
# ---- OptB_Summary
ws = sheet('OptB_Summary', 'Option B: structures compared (total return 1.75x of ticket; 24m grace; 12 quarterly payout slots)', 'S1 fixed equal instalments | S2 % of collections with floor 0.4x and cap 1.5x of the equal instalment | S3 hybrid: fixed 1.35x + share of collections, capped 2.0x | S4 fixed multiple on a coverage-gated calendar (skips quarters where cash available < 1.5x instalment).')
rows = []
NAMES = {'S1': 'S1 Fixed multiple', 'S2': 'S2 % of collections', 'S3': 'S3 Hybrid', 'S4': 'S4 Gated fixed multiple'}
for k in P:
    for T in TICKETS:
        for S in ('S1', 'S2', 'S3', 'S4'):
            o = R['B'][f'{k}|{T}|{S}']; b = o['Base']; d = o['Downside']; u = o['Upside']
            rows.append([P[k]['name'], T, NAMES[S], o['params']['pct'], b['irr'], b['moic'], b['jcost'], b['cov_min'], f"{b['cov_below']}/{b['cov_n']}", b['covr_min'], b['covr_below'], b['cov_agg'], d['irr'], d['last_pay'], d['cov_min'], d['min_cash'], u['irr'], R['B'][f'{k}|{T}|{S}']['Stress']['min_cash']])
table(ws, 4, ['Project', 'Ticket', 'Structure', 'Collections share % (S2/S3)', 'Base IRR', 'Base MOIC', 'Jalour cost of capital', 'Base min quarterly coverage (model net CF / payout)', 'Quarters < 1.5x (strict)', 'Base min coverage incl. retained-unit sale proceeds', 'Quarters < 1.5x (incl. retained)', 'Window aggregate coverage (strict)', 'Downside IRR', 'Downside last payout month', 'Downside min quarterly coverage', 'Downside peak funding gap after ticket & payouts (EGP m, negative)', 'Upside IRR', 'Stress (delivery-aligned) peak funding gap after ticket & payouts'],
      rows, [None, '0', None, '0.00%', PCT, X2, PCT, '0.00', None, '0.00', '0', '0.00', PCT, '0', '0.00', N1, PCT, N1], [14, 8, 24, 14, 10, 10, 12, 18, 12, 18, 12, 14, 12, 14, 16, 22, 10, 22])
ws.row_dimensions[4].height = 75
for i, row in enumerate(rows):
    ws.cell(5+i, 3).fill = OK if row[2].startswith('S4') else PatternFill()
r1 = 6+len(rows); ws.cell(r1, 1, 'S1 investor IRR vs total multiple (ticket 100m, equal instalments)').font = Font(name='Arial', bold=True)
rows2 = [[P[k]['name'], float(m), R['B'][f'{k}|S1_multiple_table'][m]['irr']] for k in P for m in ('1.5', '1.6', '1.7', '1.75', '1.8', '1.9', '2.0')]
table(ws, r1+1, ['Project', 'Total multiple', 'IRR'], rows2, [None, X2, PCT])
# ---- OptB_Coverage (recommended S4, per ticket)
ws = sheet('OptB_Coverage', 'Option B recommended structure (S4): payouts and coverage by quarter, base case', 'Strict = model net cash flow / payout. Incl. retained = (net cash flow + sale proceeds of the 20% retained units, ASSUMPTION: sold evenly over the 4 quarters from handover+3m, net of 4% commission) / payout.')
r = 4
for k in P:
    b = series(k); rp = retained_proceeds(k)
    for T in TICKETS:
        pay = s4_sculpted(k, T, M_B, b['net'], rp)
        ws.cell(r, 1, f"{P[k]['name']} - ticket {T}m").font = Font(name='Arial', bold=True)
        rows = []
        for m in pay_quarters(k):
            i = m//3-1; pm = float(pay[i])
            rows.append([m, float(b['net'][i]), float(rp[i]), pm, (b['net'][i]/pm) if pm > 0 else None, ((b['net'][i]+rp[i])/pm) if pm > 0 else None, 'skipped (coverage < 1.5x)' if pm == 0 else ''])
        r = table(ws, r+1, ['Month', 'Net cash flow before payout', 'Retained-unit proceeds (assumption)', 'Payout', 'Coverage strict', 'Coverage incl. retained', 'Note'], rows, ['0', N1, N1, N2, '0.00"x"', '0.00"x"', None], [10, 18, 20, 12, 14, 18, 28]) + 2
# ---- Assumptions_Log
ws = sheet('Assumptions_Log', 'Assumptions made in this analysis (all need your confirmation)')
A = [('A1', 'Investor funds 100% of ticket at t0 (GS M3, LA M12); undrawn funds earn nothing.', 'Timing of DP in model'),
 ('A2', 'Option A face value measured at model average realised price (GS avg 268/163 k EGP/m2 retail/office vs launch 200/125; LA 301/191).', 'Decision needed: launch list price would multiply the cost to Jalour by about 1.3x (GS) / 1.5x (LA)'),
 ('A3', 'Investor resells units over 4 quarters from handover+3m at face less 3% friction.', 'No data in model'),
 ('A4', 'Retained units would otherwise have been sold by Jalour on the same timing, less 4% commission.', 'Model has no retained-unit sale'),
 ('A5', 'Sold-programme source: investor units removed pro rata from every collection wave.', 'Simplification'),
 ('A6', 'Landlord waiver reduces the year-8+ excess settlement pro rata; guarantee schedule unchanged (total landlord stays above the guarantee).', 'Needs Al Ahly Sabbour'),
 ('A7', 'Pre-construction/design/permits: 10.5m (30% of the 5% professional-fee line) incurred ahead of construction. The model has no pre-construction cost.', 'Confirm real budget'),
 ('A8', 'Downside = all collections and commission shifted +12 months and price -10%; guarantee instalments and construction timing unchanged.', 'Defined by me'),
 ('A9', 'Option B total return 1.75x: chosen for IRR parity with Option A (16.4%/17.5%); not a market-tested price.', 'Decision needed'),
 ('A10', 'S4 schedule is fixed at signing from base-case cash flows and applied unchanged in downside/stress.', 'Contractual design'),
 ('A11', 'Retained-unit sale proceeds are counted as payout cash only in the "incl. retained" coverage measure.', 'Requires earmarking/sweep of those proceeds')]
table(ws, 3, ['Ref', 'Assumption', 'Note'], [list(a) for a in A], widths=[8, 140, 70])
wb.save('../01_Structuring_Analysis.xlsx'); print('saved')
