"""Extract quarterly series for Green Square (GS) and L'avenir (LA) from the source model (cached values)."""
import openpyxl, json, sys
from openpyxl.utils import column_index_from_string as CI
SRC = sys.argv[1]; OUT = sys.argv[2]
v = openpyxl.load_workbook(SRC, data_only=True)
N = 43                                   # quarters M3..M129
Q = [3*(i+1) for i in range(N)]
ncf = v['Net Cash Flow -With DP']
cod = v['Cash Out Detail']
def row(ws, r, c0=4, n=N): return [float(ws.cell(r, c0+i).value or 0) for i in range(n)]
def cod_q(r):                            # monthly cols: month m -> column m+3 ; quarter ending m = months m-2..m
    out=[]
    for m in Q:
        out.append(sum(float(cod.cell(r, mm+3).value or 0) for mm in (m-2, m-1, m))/1e6)
    return out
d = {'quarters_month': Q}
for key, (rc, rl, rcost, rcons, rcom, rsga, sheet, drow) in {
    'GS': (4, 11, 32, 28, 34, 40, 'Green Square', (35, 36)),
    'LA': (5, 12, 33, 29, 35, 41, "L'avenir", (35, 36))}.items():
    ws = v[sheet]
    # delivery-payment component of collections (sheet columns E.. = M3..), EGP thousand -> M
    deliv = [sum(float(ws.cell(r, 5+i).value or 0) for r in drow)/1000 for i in range(N)]
    d[key] = dict(collections=row(ncf, rc), landlord=row(ncf, rl), cost_total=row(ncf, rcost),
                  construction=cod_q(rcons), commission=cod_q(rcom), sga=cod_q(rsga), delivery_pay=deliv,
                  net=row(ncf, rc+42))
    # sheet-level facts
    d[key]['facts'] = {'price_comm': [ws.cell(3, c).value for c in range(5, 21)],
                       'price_admin': [ws.cell(4, c).value for c in range(5, 21)],
                       'sales_comm_pct': [ws.cell(6, c).value for c in range(5, 21)],
                       'sales_admin_pct': [ws.cell(7, c).value for c in range(5, 21)],
                       'total_sales_80_k': ws['D26'].value, 'sales_comm_80_k': ws['D24'].value, 'sales_admin_80_k': ws['D25'].value,
                       'area_comm': ws['B6'].value, 'area_admin_sales_basis': ws['B7'].value}
json.dump(d, open(OUT, 'w'), indent=1)
for k in ('GS', 'LA'):
    x = d[k]; print(k, 'coll', round(sum(x['collections']), 2), 'landlord', round(sum(x['landlord']), 2), 'cost', round(sum(x['cost_total']), 2),
        'cons+com+sga', round(sum(x['construction'])+sum(x['commission'])+sum(x['sga']), 2), 'deliv', round(sum(x['delivery_pay']), 2), 'net', round(sum(x['net']), 2))
