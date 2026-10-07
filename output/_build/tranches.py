"""Rebuild per-tranche collections from the source model sheets (cached values) and verify against row 70."""
import openpyxl, json, sys
import numpy as np
SRC = sys.argv[1]; OUT = sys.argv[2]
v = openpyxl.load_workbook(SRC, data_only=True); f = openpyxl.load_workbook(SRC)
NQ = 60
def g(ws, r, c): return float(ws.cell(r, c).value or 0)
res = {}
for key, sh in (('GS', 'Green Square'), ('LA', "L'avenir")):
    ws = v[sh]; wf = f[sh]
    # delivery quarter column: the column where row 35 first non-zero
    dcol = [c for c in range(5, 40) if g(ws, 35, c) != 0][0]
    print(key, 'delivery payment column', openpyxl.utils.get_column_letter(dcol), ws.cell(2, dcol).value, wf.cell(35, dcol).value)
    tr = []
    for j in range(16):                       # tranches in columns E..T
        c = 5+j; Sc, Sa = g(ws, 24, c), g(ws, 25, c)
        if Sc+Sa == 0: continue
        flow = np.zeros(NQ+2); fc = np.zeros(NQ+2); fa = np.zeros(NQ+2)
        def add(col, val, vc=0.0, va=0.0):
            flow[col-5] += val; fc[col-5] += vc; fa[col-5] += va
        add(c, g(ws, 31, c)+g(ws, 32, c), g(ws, 31, c), g(ws, 32, c))                      # down payment
        add(c+1, g(ws, 33, c+1)+g(ws, 34, c+1), g(ws, 33, c+1), g(ws, 34, c+1))               # 3-month payment (copy of DP next quarter)
        rr = 38+2*j
        for cc in range(5, 5+NQ): add(cc, g(ws, rr, cc)+g(ws, rr+1, cc), g(ws, rr, cc), g(ws, rr+1, cc))   # instalments (comm+admin)
        add(max(c, dcol), Sc*g(ws, 12, c)/100 + Sa*g(ws, 12, c)/100, Sc*g(ws, 12, c)/100, Sa*g(ws, 12, c)/100)     # delivery payment
        tr.append(dict(col=c, month=3*(c-4), sales_comm=Sc, sales_admin=Sa, price_comm=g(ws, 3, c), price_admin=g(ws, 4, c),
                       area_comm=Sc/g(ws, 3, c) if g(ws, 3, c) else 0, area_admin=Sa/g(ws, 4, c) if g(ws, 4, c) else 0, flow=(flow[:NQ]/1000).tolist(), flow_comm=(fc[:NQ]/1000).tolist(), flow_admin=(fa[:NQ]/1000).tolist()))
    tot = sum(np.array(t['flow']) for t in tr)
    model = np.array([g(ws, 70, 5+i)/1000 for i in range(NQ)])
    print(key, 'tranches', len(tr), 'sum recon', round(tot.sum(), 3), 'model', round(model.sum(), 3), 'max abs diff', round(np.abs(tot-model).max(), 4))
    res[key] = dict(tranches=tr, delivery_col=dcol)
json.dump(res, open(OUT, 'w'))
