import subprocess, re, os, tempfile
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY = RGBColor(0x17, 0x32, 0x4D); BRASS = RGBColor(0xB8, 0x89, 0x3B); GREY = RGBColor(0x6B, 0x72, 0x80); INK = RGBColor(0x1B, 0x1F, 0x2A)
def fmt(x, d=1):
    return 'n/a' if x is None or x != x else f'{x:,.{d}f}'
def pct(x, d=1): return 'n/a' if x is None or x != x else f'{x*100:.{d}f}%'
def mx(x, d=2): return f'{x:.{d}f}x'
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)
def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)
def set_borders(table, color='C9CFD8', sz=4):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color); b.append(e)
    tblPr.append(b)
def add_field(par, instr):
    r = par.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r3._r.append(f2)
class Doc:
    def __init__(self, title, header_text, author='Jalour Developments'):
        self.d = Document(); d = self.d
        sec = d.sections[0]; sec.page_width = Cm(21.0); sec.page_height = Cm(29.7); sec.left_margin = sec.right_margin = Cm(2.3); sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)
        st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri'); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.12
        for n, sz, sp in (('Heading 1', 17, 18), ('Heading 2', 13, 12), ('Heading 3', 11, 8)):
            h = d.styles[n]; h.font.name = 'Cambria'; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = NAVY; h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria'); h.element.rPr.rFonts.set(qn('w:ascii'), 'Cambria'); h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Cambria')
            h.paragraph_format.space_before = Pt(sp); h.paragraph_format.space_after = Pt(6); h.paragraph_format.keep_with_next = True
        cp = d.core_properties; cp.title = title; cp.author = author; cp.subject = 'Strictly private and confidential'; cp.keywords = ''; cp.comments = ''; cp.last_modified_by = author
        self.header_text = header_text; self.headings = []; self.fig_n = 0; self.tab_n = 0
        hp = sec.header.paragraphs[0]; hp.text = header_text; hp.runs[0].font.size = Pt(8.5); hp.runs[0].font.color.rgb = GREY
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = fp.add_run('Page '); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        add_field(fp, 'PAGE')
        for rr in fp.runs: rr.font.size = Pt(8.5); rr.font.color.rgb = GREY
        sec.different_first_page_header_footer = True
    def h1(self, text, page_break=True):
        p = self.d.add_heading(text, 1); self.headings.append((1, text))
        if page_break: p.paragraph_format.page_break_before = True
        return p
    def h2(self, text): p = self.d.add_heading(text, 2); self.headings.append((2, text)); return p
    def h3(self, text): return self.d.add_heading(text, 3)
    def p(self, text, bold=False, italic=False, size=None, color=None, align=None, after=None, keep=False):
        par = self.d.add_paragraph(); self._runs(par, text, bold, italic, size, color)
        if align == 'center': par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None: par.paragraph_format.space_after = Pt(after)
        if keep: par.paragraph_format.keep_with_next = True
        return par
    def _runs(self, par, text, bold=False, italic=False, size=None, color=None):
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if not part: continue
            b = part.startswith('**') and part.endswith('**'); t = part[2:-2] if b else part
            r = par.add_run(t); r.bold = bold or b; r.italic = italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
    def bullets(self, items, style='List Bullet'):
        for t in items:
            par = self.d.add_paragraph(style=style); self._runs(par, t); par.paragraph_format.space_after = Pt(3)
    def numbered(self, items): self.bullets(items, 'List Number')
    def table(self, rows, widths_cm, caption=None, align_right_from=1, font=9.5, header=True, bold_first_col=False, total_row=False):
        d = self.d
        if caption:
            self.tab_n += 1; cp = d.add_paragraph(); r = cp.add_run(f'Table {self.tab_n}: {caption}'); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = NAVY; cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(3)
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; set_borders(t); cell_margins(t)
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
            if ri == 0 and header:
                th = OxmlElement('w:tblHeader'); trPr.append(th)
            for ci, val in enumerate(row):
                c = tr.cells[ci]; c.width = Cm(widths_cm[ci]); c.text = ''
                par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
                r = par.add_run(str(val)); r.font.size = Pt(font)
                if ri == 0 and header: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, '17324D')
                else:
                    if (bold_first_col and ci == 0) or (total_row and ri == len(rows)-1): r.bold = True
                    if ri % 2 == 0: shade(c, 'F2F4F7')
                if ci >= align_right_from and align_right_from >= 0: par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if ri == 0 and ci >= align_right_from and align_right_from >= 0: par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(4); return t
    def figure(self, path, caption, source, width_cm=15.5):
        self.fig_n += 1; d = self.d
        par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.keep_with_next = True; par.add_run().add_picture(path, width=Cm(width_cm))
        c = d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = c.add_run(f'Figure {self.fig_n}: {caption}. '); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = NAVY
        r2 = c.add_run(f'Source: {source}'); r2.italic = True; r2.font.size = Pt(9); r2.font.color.rgb = GREY
    def callout(self, text, fill='FBF3E4'):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; cell_margins(t, 100, 100, 160, 160)
        c = t.rows[0].cells[0]; c.width = Cm(16.4); shade(c, fill); c.text = ''; par = c.paragraphs[0]; self._runs(par, text, size=10); par.paragraph_format.space_after = Pt(0)
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
    def save(self, path): self.d.save(path)
def pdf_pages(docx_path):
    out = tempfile.mkdtemp(); subprocess.run(['python3', '/root/.claude/skills/synced/8e697198-15eb-43a6-b507-8c3757e6e339_12c2af44-09f8-4a5d-8395-d9d914cc879e/docx/scripts/office/soffice.py', '--headless', '--convert-to', 'pdf', '--outdir', out, docx_path], capture_output=True)
    pdf = os.path.join(out, os.path.basename(docx_path).replace('.docx', '.pdf'))
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout).group(1))
    pages = [subprocess.run(['pdftotext', '-f', str(i), '-l', str(i), '-layout', pdf, '-'], capture_output=True, text=True).stdout for i in range(1, n+1)]
    return pdf, pages
def find_pages(headings, pages):
    res = {}
    for lvl, h in headings:
        if lvl != 1: continue
        for i, tx in enumerate(pages):
            if i < 2: continue
            lines = [l.strip() for l in tx.splitlines()]
            if any(l == h for l in lines): res[h] = i+1; break
    return res
