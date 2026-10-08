import sys, os, re, json, zipfile, glob; sys.path.insert(0, '.')
import openpyxl
import packdata as PD
from qa import docx_text, pptx_text, variants, indep_irr
OUT = '/home/user/Jalour_inv_pitch/output'
def run(Te=62.5):
    folder = f'{OUT}/COMBINED_125M'; work = '/tmp/w/work_comb'; rep = ['# QA report: Green Square and L\'avenir combined pack, ticket EGP 125m (62.5 + 62.5)\n']; fails = 0
    C = PD.build_combined(Te); mp = json.load(open(f'{work}/annex_map.json'))['M']; G, L = C['GS'], C['LA']; S1 = C['B']['S1']; sp = C['sponsor']; Sc = sp['scen']
    wb = openpyxl.load_workbook(f'{folder}/Financial_Annex.xlsx', data_only=True)
    def cell(ref): sh, a = ref.split('!'); return wb[sh.strip("'")][a.replace('$', '')].value
    S = {k: cell(r) for k, r in mp.items() if not k.startswith(('scen|', 'sens|', 'cp|'))}
    texts = {'Deck': pptx_text(f'{folder}/Deck.pptx'), 'Memo': docx_text(f'{folder}/Investment_Memo.docx'), 'DD': docx_text(f'{folder}/DD_Memo.docx'), 'NDA': docx_text(f'{folder}/NDA_Template.docx')}
    spec = [('Combined ticket', 'tick', 'm0', 'Deck Memo DD'), ('Option A face value (combined)', 'face', 'm0', 'Deck Memo'), ('Option A value at delivery list', 'val', 'm0', 'Memo'), ('Option A investor IRR base', 'a_irr', 'pct1', 'Deck Memo'), ('Option A MOIC base', 'a_moic', 'x2', 'Deck'), ('Option A IRR downside', 'a_dirr', 'pct1', 'Deck Memo'), ('Option A IRR upside', 'a_uirr', 'pct1', 'Deck Memo'),
            ('Jalour cost of capital, Option A', 'j_irr', 'pct1', 'Memo'), ('Option B total return', 'b_total', 'm0', 'Deck Memo'), ('Option B investor IRR', 'b_irr', 'pct1', 'Deck Memo'), ('Option B MOIC', 'b_moic', 'x2', 'Deck Memo'), ('Combined lowest position, base', 'cp_min', 'm1', 'Deck Memo DD')]
    rep.append('## 1. Number tie-out (annex vs deck, memo, DD memo)\n\n| Figure | Annex cell | Value | Found in | Result |\n|---|---|---|---|---|'); miss = 0
    for lab, key, kind, docs in spec:
        val = S[key]; v = abs(val) if key == 'cp_min' else val; vs = variants(kind, v); res = {dn: any(s in texts[dn] for s in vs) for dn in docs.split()}; ok = all(res.values()); miss += (not ok)
        rep.append(f'| {lab} | {mp[key].replace("$", "")} | {val:.4f} | ' + ', '.join(f'{a}: {"yes" if b else "NO"}' for a, b in res.items()) + f' | {"OK" if ok else "CHECK"} |')
    for lab, val, kind, docs in [('GS base peak shortage', G['fund']['model_peak'], 'm1', 'Deck Memo'), ('LA base peak shortage', L['fund']['model_peak'], 'm1', 'Deck Memo'), ('GS net cash flow', G['base']['net'], 'm1', 'Deck Memo'), ('LA net cash flow', L['base']['net'], 'm1', 'Deck Memo'), ('GS NPV', G['base']['npv'], 'm1', 'Deck Memo'), ('LA NPV', L['base']['npv'], 'm1', 'Deck Memo'),
                                 ('Combined peak, handover payments', -Sc['Delivery payments at handover']['comb_peak'], 'm0', 'Deck Memo DD'), ('Combined peak, downside', -Sc['Downside: sales +12 months, prices -10%']['comb_peak'], 'm0', 'Deck Memo DD'), ('Combined low month', sp['comb_min_m'], 'int', 'Deck Memo'), ('Combined positive from month', sp['comb_pos_month'], 'int', 'Deck Memo'),
                                 ('Instalment per project', G['B']['mult']*Te/12, 'm1', 'Memo'), ('Option B per project', G['B']['mult']*Te, 'm0', 'Deck Memo')]:
        vs = variants(kind, val) | ({f'{val:,.2f}'} if 'Instalment' in lab else set()) | ({f'{val:,.2f}'} if 'per project' in lab else set()); res = {dn: any(s in texts[dn] for s in vs) for dn in docs.split()}; ok = all(res.values()); miss += (not ok)
        rep.append(f'| {lab} | python | {val:.4f} | ' + ', '.join(f'{a}: {"yes" if b else "NO"}' for a, b in res.items()) + f' | {"OK" if ok else "CHECK"} |')
    rep.append(f'\nTie-out misses: {miss}\n'); fails += miss
    # 2 independent
    t = [(3*(i+1)-PD.OFF['GS'])/12 for i in range(60)]; tk = {k_: [(3*(i+1)-PD.OFF[k_])/12 for i in range(60)] for k_ in ('GS', 'LA')}
    def rv(sheet, r): ws = wb[sheet]; return [ws.cell(r, 4+i).value or 0 for i in range(60)]
    add = lambda a, b: [x+y for x, y in zip(a, b)]
    invA = add(rv('GS_Option_A', 17), rv('LA_Option_A', 17)); jal = add(rv('GS_Option_A', 21), rv('LA_Option_A', 21)); invB = add(rv('GS_Option_B', 14), rv('LA_Option_B', 14))
    chk = [('Option A combined investor IRR', indep_irr(invA, t), S['a_irr'], 1e-4), ('Option A combined MOIC', sum(x for x in invA if x > 0)/-sum(x for x in invA if x < 0), S['a_moic'], 1e-4), ('Option A Jalour cost of capital', indep_irr(jal, t, 0.0), S['j_irr'], 1e-4), ('Option A Jalour NPV @14%', sum(c/(1.14**ti) for c, ti in zip(jal, t)), S['j_npv'], 1e-4), ('Option B combined investor IRR', indep_irr(invB, t), S['b_irr'], 1e-4)]
    for p, D in (('GS', G), ('LA', L)):
        net = rv(f'{p}_Project_CF', 13); chk.append((f'{p} project NPV @14% (own timeline)', sum(c/(1.14**ti) for c, ti in zip(net, tk[p])), D['base']['npv'], 1e-4)); chk.append((f'{p} net cash flow', sum(net), D['base']['net'], 1e-4))
    nets = add(rv('GS_Project_CF', 13), rv('LA_Project_CF', 13)); cum = 0; mn = 0; i_mn = 0
    for i, x in enumerate(nets):
        cum += x
        if cum < mn: mn, i_mn = cum, i
    chk.append(('Combined lowest cumulative position', mn, -S['cp_min'] if S['cp_min'] > 0 else S['cp_min'], 1e-4))
    rep.append('## 2. Arithmetic check: independent pure-Python recomputation vs spreadsheet\n\n| Measure | Independent | Spreadsheet / engine | Abs difference | Result |\n|---|---|---|---|---|'); bad = 0
    for lab, a, b, tol in chk:
        diff = abs(a-b); ok = diff <= tol; bad += (not ok); rep.append(f'| {lab} | {a:.6f} | {b:.6f} | {diff:.2e} | {"OK" if ok else "EXPLAIN"} |')
    rc = wb['Combined_Reconciliation']; rep.append(f'\nAnnex reconciliation sheet final check: {rc.cell(rc.max_row, 4).value}\n'); fails += bad
    # 3 disclosure scan
    pats = ['retain', 'AT' + ' EAST', 'At' + ' East', 'AT' + ' East', 'total raise', 'total programme', 'other investor', 'second investor', 'only external investor', 'last external investor', 'entire balance sheet', 'own balance sheet', 'combined raise', 'exclusive investor', 'sole investor', 'ring-fenced for']
    rep.append('## 3. Disclosure scan (every part of every file)\n\nBoth project names are permitted in this pack: it is the single combined investor pack.\n\n| File | Parts scanned | Prohibited hits | Confidentiality statement |\n|---|---|---|---|'); ht = 0
    for f in sorted(glob.glob(folder + '/*')):
        parts = 0; hits = []; conf = False
        if f.endswith('.pdf'):
            import subprocess; tx = subprocess.run(['pdftotext', f, '-'], capture_output=True, text=True).stdout; parts = 1
            hits = [f'{p_} in pdf text' for p_ in pats if p_.lower() in tx.lower()]; conf = bool(re.search(r'(?i)confidential', tx))
            ht += len(hits) + (not conf); rep.append(f'| {os.path.basename(f)} | {parts} | {len(hits)} {"; ".join(hits[:5])} | {"yes" if conf else "MISSING"} |'); continue
        z = zipfile.ZipFile(f)
        for nme in z.namelist():
            if not nme.endswith(('.xml', '.rels', '.txt', '.json')): continue
            tx = z.read(nme).decode('utf8', 'ignore'); parts += 1; clean = re.sub(r'<[^>]+>', ' ', tx)
            for p_ in pats:
                if p_.lower() in clean.lower(): hits.append(f'{p_} in {nme}')
            if re.search(r'(?i)confidential', clean): conf = True
        if f.endswith('.xlsx'):
            wsx = openpyxl.load_workbook(f); hidden = [w.title for w in wsx if w.sheet_state != 'visible']; cm = [(w.title, c.coordinate) for w in wsx for r in w.iter_rows() for c in r if c.comment]
            if hidden or cm: hits.append(f'hidden sheets {hidden} comments {cm}')
        ht += len(hits) + (not conf); rep.append(f'| {os.path.basename(f)} | {parts} | {len(hits)} {"; ".join(hits[:5])} | {"yes" if conf else "MISSING"} |')
    # required wording
    req = [('Neutral funding wording', 'The balance of project funding is provided by Jalour sponsor equity, project collections and other capital sources', ('Deck', 'Memo', 'DD')), ('Further capital clause', 'Jalour may raise further capital at project or holding level', ('Deck', 'Memo', 'DD')), ('No cross-collateral', 'cross-collateral', ('Deck', 'Memo', 'DD')), ('Confidential banner', 'CONFIDENTIAL', ('Deck', 'Memo', 'DD', 'NDA')), ('NDA covers all information', 'confidential', ('NDA',))]
    for lab, s, docs in req:
        for dn in docs:
            ok = s.lower() in re.sub(r'\s+', ' ', texts[dn]).lower(); ht += (not ok); rep.append(f'- {lab} in {dn}: {"yes" if ok else "MISSING"}')
    for f in sorted(glob.glob(folder + '/*')):
        for p_ in pats:
            if p_.lower() in os.path.basename(f).lower(): ht += 1; rep.append(f'FILE NAME HIT: {f}')
    rep.append(f'\nDisclosure and wording failures: {ht}\n'); fails += ht
    # NDA content
    nd = re.sub(r'\s+', ' ', texts['NDA']).lower()
    for lab, s in [('Green Square named', 'green square'), ("L'avenir named", "l'avenir"), ('Covers financial information of both projects', 'financial'), ('Non-circumvention', 'circumvent'), ('Return or destroy', 'destr'), ('Term', 'term'), ('Governing law', 'governing law'), ('No obligation / no exclusivity', 'no obligation')]:
        ok = s in nd; fails += (not ok); rep.append(f'- NDA template, {lab}: {"present" if ok else "MISSING"}')
    rep.append(f'\n**Total failures requiring attention: {fails}**')
    open(f'{OUT}/QA_Reports/QA_COMBINED_125M.md', 'w').write('\n'.join(rep)); return fails
if __name__ == '__main__': print('failures', run())
